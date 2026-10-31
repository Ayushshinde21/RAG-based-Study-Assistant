"""Text extraction for supported study documents.

Supported inputs: PDF, DOCX, PPTX, TXT and Markdown.
Page and slide markers are retained so the chunker can attach source locations
as chunk metadata.
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from zipfile import BadZipFile, ZipFile


SUPPORTED_DOCUMENT_EXTENSIONS = {"pdf", "docx", "pptx", "txt", "md"}
MAX_OFFICE_ARCHIVE_ENTRIES = 10_000
MAX_OFFICE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_OFFICE_MEMBER_BYTES = 50 * 1024 * 1024


class DocumentReadError(RuntimeError):
    """Raised when a document cannot be converted into usable text."""


def clean_document_text(text: str) -> str:
    """Normalize extraction artefacts without flattening meaningful structure."""
    if not isinstance(text, str):
        raise TypeError("Document text must be a string.")

    text = unicodedata.normalize("NFC", text)
    text = text.replace("\ufeff", "").replace("\x00", "")
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()




def _validate_office_archive(path: Path) -> None:
    """Reject malformed or oversized DOCX/PPTX ZIP containers before parsing."""
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
    except (BadZipFile, OSError, ValueError) as exc:
        raise DocumentReadError("The Office document is not a valid DOCX/PPTX archive.") from exc

    if len(entries) > MAX_OFFICE_ARCHIVE_ENTRIES:
        raise DocumentReadError("The Office document contains too many archive entries.")

    total_uncompressed = 0
    for entry in entries:
        if entry.file_size < 0 or entry.compress_size < 0:
            raise DocumentReadError("The Office document contains invalid archive metadata.")
        if entry.file_size > MAX_OFFICE_MEMBER_BYTES:
            raise DocumentReadError("The Office document contains an oversized embedded file.")
        total_uncompressed += entry.file_size
        if total_uncompressed > MAX_OFFICE_UNCOMPRESSED_BYTES:
            raise DocumentReadError("The Office document expands beyond the allowed processing size.")


def _read_docx(path: Path) -> str:
    try:
        from docx import Document as WordDocument

        document = WordDocument(str(path))
    except Exception as exc:
        raise DocumentReadError(f"Could not read DOCX document: {exc}") from exc

    pieces: list[str] = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if value:
            pieces.append(value)

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(values):
                pieces.append(" | ".join(values))

    return "\n\n".join(pieces)


def _read_pptx(path: Path) -> str:
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
    except Exception as exc:
        raise DocumentReadError(f"Could not read PowerPoint document: {exc}") from exc

    slides: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        pieces: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                pieces.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(values):
                        pieces.append(" | ".join(values))
        if pieces:
            slides.append(f"[Slide {slide_number}]\n" + "\n".join(pieces))

    return "\n\n".join(slides)


def read_document_text(path: str, filename: str | None = None) -> str:
    """Extract normalized text while preserving PDF page and PPTX slide markers."""
    document_path = Path(path)
    if not document_path.exists() or not document_path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    display_name = filename or document_path.name
    extension = Path(display_name.replace("\\", "/")).suffix.lower().lstrip(".")
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        accepted = ", ".join(sorted(SUPPORTED_DOCUMENT_EXTENSIONS))
        raise DocumentReadError(f"Unsupported document format. Supported formats: {accepted}.")

    if extension == "pdf":
        try:
            from utils.pdf_reader import read_pdf_text

            extracted = read_pdf_text(str(document_path), with_page_markers=True)
        except Exception as exc:
            message = str(exc).strip() or exc.__class__.__name__
            raise DocumentReadError(message) from exc
    elif extension in {"txt", "md"}:
        try:
            extracted = document_path.read_text(encoding="utf-8-sig", errors="replace")
        except OSError as exc:
            raise DocumentReadError(f"Could not read text document: {exc}") from exc
    elif extension == "docx":
        _validate_office_archive(document_path)
        extracted = _read_docx(document_path)
    else:
        _validate_office_archive(document_path)
        extracted = _read_pptx(document_path)

    cleaned = clean_document_text(extracted)
    if not cleaned:
        if extension == "pdf":
            raise DocumentReadError(
                "No selectable text was found in this PDF. It may be scanned or image-only; "
                "OCR is required before it can be indexed."
            )
        raise DocumentReadError("No readable text was found in this document.")

    return cleaned
