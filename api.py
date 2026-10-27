"""
Campus Study Desk - FastAPI backend.

Put this file in the project root, next to main.py / core / utils / features,
so the existing imports and the chroma_db/ and transcripts/ folders keep working.

Run locally:   uvicorn api:app --reload --port 8000
API docs:      http://127.0.0.1:8000/docs
"""
import gc
import json
import os
import re
import tempfile
import threading
import time
import uuid
from contextlib import asynccontextmanager
from datetime import datetime
from typing import Literal
from urllib.parse import parse_qs, urlparse

from dotenv import load_dotenv

load_dotenv()

from fastapi import Depends, FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

import auth_store

ALLOWED_EXTENSIONS = {"mp4", "mkv", "avi", "mov", "webm", "mp3", "wav", "m4a", "ogg", "flac"}
REFERENCE_EXTENSIONS = {"pdf", "docx", "pptx", "txt", "md"}
# Cloud Run rejects HTTP/1 request bodies over 32 MiB, so the hosted default is 30.
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "200"))
MAX_REFERENCE_FILE_MB = int(os.getenv("MAX_REFERENCE_FILE_MB", "15"))
MAX_REFERENCE_FILES = 5
MAX_REFERENCE_TEXT_CHARS = 250_000
YT_HOSTS = {"youtube.com", "www.youtube.com", "m.youtube.com", "music.youtube.com", "youtu.be", "www.youtu.be"}
TITLES_PATH = "lecture_titles.json"
# YouTube blocks downloads from cloud servers (HTTP 403), so the hosted build sets this to 0.
YOUTUBE_ENABLED = os.getenv("YOUTUBE_ENABLED", "1") == "1"

@asynccontextmanager
async def lifespan(_app):
    # On Cloud Run (WARMUP=1) load every model before the first request,
    # so the first user doesn't wait for them. Skipped locally for fast reloads.
    if os.getenv("WARMUP") == "1":
        import utils.transcriber  # noqa: F401  (loads Whisper)
        import core.reranker      # noqa: F401  (loads the CrossEncoder)
        from core.vector_store import get_embeddings
        get_embeddings()
    yield


app = FastAPI(title="Campus Study Desk API", lifespan=lifespan)

origins = [o.strip() for o in os.getenv(
    "FRONTEND_ORIGINS", "http://127.0.0.1:5500,http://localhost:5500"
).split(",") if o.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_methods=["GET", "POST", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type", "X-CSRF-Token"],
    allow_credentials=True,
)

SESSION_COOKIE = "campus_session"
SESSION_COOKIE_SECURE = os.getenv("SESSION_COOKIE_SECURE", "0") == "1"
LOGIN_ATTEMPTS: dict[str, list[float]] = {}
LOGIN_RATE_WINDOW_SECONDS = 300
LOGIN_RATE_LIMIT = 12


def current_user(request: Request) -> dict:
    session = auth_store.get_session(request.cookies.get(SESSION_COOKIE))
    if not session:
        raise HTTPException(status_code=401, detail="Please sign in to continue.")
    request.state.csrf_token = session["csrf_token"]
    return session["user"]


def require_roles(*roles: str):
    def dependency(user: dict = Depends(current_user)) -> dict:
        if user["role"] not in roles:
            raise HTTPException(status_code=403, detail="Your account does not have permission to do that.")
        return user
    return dependency


STUDENT_OR_ADMIN = require_roles("student", "admin")
FARMER_ONLY = require_roles("farmer")
ADMIN_ONLY = require_roles("admin")


@app.middleware("http")
async def csrf_guard(request: Request, call_next):
    if request.method in {"POST", "PATCH", "DELETE", "PUT"} and request.url.path.startswith("/api/"):
        public_auth_routes = {"/api/auth/login", "/api/auth/register"}
        if request.url.path not in public_auth_routes and request.cookies.get(SESSION_COOKIE):
            session = auth_store.get_session(request.cookies.get(SESSION_COOKIE))
            supplied = request.headers.get("x-csrf-token", "")
            if not session or not supplied or not secrets_compare(supplied, session["csrf_token"]):
                return JSONResponse(status_code=403, content={"detail": "Your security token is missing or expired. Refresh the page and try again."})
    return await call_next(request)


def secrets_compare(left: str, right: str) -> bool:
    import hmac
    return hmac.compare_digest(left, right)


def set_session_cookie(response: Response, token: str) -> None:
    response.set_cookie(
        SESSION_COOKIE, token, max_age=auth_store.SESSION_TTL_SECONDS,
        httponly=True, secure=SESSION_COOKIE_SECURE, samesite="strict", path="/",
    )


def clear_session_cookie(response: Response) -> None:
    response.delete_cookie(SESSION_COOKIE, path="/", httponly=True,
                           secure=SESSION_COOKIE_SECURE, samesite="strict")


def lecture_belongs_to_user(lecture: dict, user: dict) -> bool:
    if user["role"] == "admin":
        return True
    if user["role"] != "student":
        return False
    try:
        return lecture.get("owner_user_id") is not None and int(lecture["owner_user_id"]) == int(user["id"])
    except (TypeError, ValueError):
        return False


def require_lecture_access(lecture_id: str, user: dict) -> dict:
    from core.vector_store import get_lecture
    lecture = get_lecture(lecture_id)
    if not lecture or not lecture_belongs_to_user(lecture, user):
        # Use 404 rather than confirming that another user's lecture exists.
        raise HTTPException(404, "Lecture not found")
    return lecture

# ── In-memory state (replaces st.session_state) ──────────────────────────────
JOBS: dict[str, dict] = {}                 # job_id -> progress of a processing run
RUNTIME: dict[str, dict] = {}              # lecture_id -> {docs, dense, bm25} (shared retrievers)
ENGINES: dict[tuple, object] = {}          # (lecture_id, session_id) -> RAGEngine (own chat memory)
SUMMARIES: dict[str, str] = {}
DETAILED_NOTES: dict[str, str] = {}
PROCESS_LOCK = threading.Semaphore(1)      # one heavy pipeline at a time (protects RAM)
LOAD_LOCK = threading.Lock()
TITLE_LOCK = threading.Lock()
MAX_ENGINES = 200


# ── Helpers ──────────────────────────────────────────────────────────────────
def redact(text) -> str:
    text = str(text)
    for name, value in os.environ.items():
        if value and len(value) >= 8 and any(t in name.upper() for t in ("KEY", "TOKEN", "SECRET", "PASSWORD")):
            text = text.replace(value, "[hidden]")
    return text


def friendly_error(exc: Exception) -> str:
    msg = redact(exc)
    low = msg.lower()
    if "403" in msg or "forbidden" in low:
        return "The download was refused (HTTP 403). Try another video, or upload the recording as a file."
    if "api key" in low or "api_key" in low or "401" in msg or "unauthorized" in low:
        return "The AI service rejected the request. Check the API keys on the server."
    if "rate limit" in low or "429" in msg or "quota" in low:
        return "The AI service rate limit was reached. Please wait a moment and try again."
    return msg or exc.__class__.__name__


def valid_youtube_url(url: str) -> bool:
    url = (url or "").strip()
    if not url or len(url) > 500 or re.search(r"\s", url):
        return False
    try:
        p = urlparse(url)
        host = (p.hostname or "").lower()
    except ValueError:
        return False
    if p.scheme not in ("http", "https") or host not in YT_HOSTS:
        return False
    if host.endswith("youtu.be"):
        return bool(p.path.strip("/"))
    return bool(parse_qs(p.query).get("v")) or p.path.startswith(("/shorts/", "/embed/", "/live/"))


def load_titles() -> dict:
    try:
        with open(TITLES_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return {}


def save_titles(titles: dict):
    with open(TITLES_PATH, "w", encoding="utf-8") as f:
        json.dump(titles, f, indent=2)


def with_title(lecture: dict) -> dict:
    out = dict(lecture)
    out["title"] = load_titles().get(lecture["id"], lecture.get("title", "Untitled lecture"))
    return out


def get_runtime(lecture_id: str) -> dict:
    """Load a saved lecture's retrievers on first use (same as open_lecture in app.py)."""
    rt = RUNTIME.get(lecture_id)
    if rt:
        return rt
    from core.vector_store import get_lecture, load_vector_store, get_all_docs_from_store, get_dense_retriever
    from core.bm25_index import build_bm25_retriever

    lecture = get_lecture(lecture_id)
    if not lecture:
        raise HTTPException(404, "Lecture not found")
    with LOAD_LOCK:
        rt = RUNTIME.get(lecture_id)
        if rt:
            return rt
        try:
            vs = load_vector_store(lecture["collection_name"])
            docs = get_all_docs_from_store(vs)
            if not docs:
                raise RuntimeError("The saved index for this lecture is empty or missing.")
            rt = {"docs": docs, "dense": get_dense_retriever(vs, k=10), "bm25": build_bm25_retriever(docs, k=10)}
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(500, f"Couldn't open this lecture. {friendly_error(e)}")
        RUNTIME[lecture_id] = rt
    return rt


def get_engine(lecture_id: str, session_id: str):
    key = (lecture_id, session_id)
    eng = ENGINES.get(key)
    if eng is None:
        from core.rag_engine import RAGEngine
        rt = get_runtime(lecture_id)
        eng = RAGEngine(rt["dense"], rt["bm25"])
        if len(ENGINES) >= MAX_ENGINES:
            ENGINES.pop(next(iter(ENGINES)))
        ENGINES[key] = eng
    return eng


# ── Processing pipeline (process_uploaded_file from app.py, minus Streamlit) ─
STEPS = ["Extracting audio", "Transcribing lecture", "Chunking transcript",
         "Building search index", "Warming up reranker", "Saving to library"]


def set_step(job: dict, i: int):
    job["step"], job["stage"] = i + 1, STEPS[i]


def run_job(job_id: str, source: str, is_url: bool, owner_user_id: int, reference_files: list[dict] | None = None):
    job = JOBS[job_id]
    reference_files = reference_files or []
    job.setdefault("warnings", [])
    audio_path = None
    vector_store = None
    lecture_id = None
    job["status"] = "queued"
    try:
        with PROCESS_LOCK:
            job["status"] = "running"

            video_title = None
            transcript_path = None
            set_step(job, 0)

            # Captions are far cheaper than speech recognition for multi-hour YouTube videos.
            # If English captions are unavailable, continue with the normal audio pipeline.
            if is_url:
                job["stage"] = "Checking for YouTube captions"
                try:
                    from utils.audio_processor import download_youtube_captions
                    caption_result = download_youtube_captions(source)
                except Exception as e:
                    print(f"[job {job_id}] caption-first processing skipped: {friendly_error(e)}")
                    caption_result = None

                if caption_result:
                    from utils.transcriber import save_transcript
                    segments, video_title = caption_result
                    safe_title = re.sub(r"[^A-Za-z0-9._-]+", "_", video_title).strip("._")[:150] or "YouTube_lecture"
                    transcript_path = save_transcript(segments, os.path.join("audio_files", safe_title + ".mp3"))
                    del segments
                    gc.collect()
                    job["step"] = 2
                    job["stage"] = "YouTube captions ready"

            if transcript_path is None:
                if is_url and not YOUTUBE_ENABLED:
                    raise RuntimeError(
                        "No usable English YouTube captions were available. This server has YouTube audio downloading disabled. "
                        "Try a video with English captions or run the project locally with YOUTUBE_ENABLED=1."
                    )
                set_step(job, 0)
                from utils.audio_processor import get_audio_path
                audio_path = get_audio_path(source)
                if not audio_path or not os.path.exists(audio_path):
                    raise RuntimeError("No audio could be extracted from this source.")
                size = os.path.getsize(audio_path)
                print(f"[job {job_id}] audio file: {audio_path} ({size / 1024:.0f} KB)")
                if size < 2048:
                    raise RuntimeError(
                        f"The audio file is nearly empty ({size} bytes). YouTube probably "
                        "cut the download. Try again, try another video, or upload the file."
                    )

                set_step(job, 1)
                from utils.transcriber import transcribe_with_timestamps, save_transcript
                def transcription_progress(message: str):
                    job["stage"] = message
                segments = transcribe_with_timestamps(audio_path, progress_callback=transcription_progress)
                if not segments:
                    raise RuntimeError("The transcription came back empty. The recording may be silent.")
                transcript_path = save_transcript(segments, audio_path)
                del segments
                gc.collect()

            set_step(job, 2)
            from core.chunker import chunk_from_file
            docs = chunk_from_file(transcript_path, method="semantic")
            if not docs:
                raise RuntimeError("No text chunks could be created from the transcript.")
            for doc in docs:
                doc.metadata["source_type"] = "lecture_transcript"

            # Lecturer-provided material is optional. A malformed attachment should
            # not prevent a valid lecture recording from being processed.
            if reference_files:
                from core.chunker import chunk_text
                for index, item in enumerate(reference_files, start=1):
                    try:
                        job["stage"] = f"Reading lecturer material {index} of {len(reference_files)}"
                        text = extract_reference_text(item["path"], item["filename"])
                        if not text.strip():
                            raise ValueError("No readable text was found in this file.")
                        if len(text) > MAX_REFERENCE_TEXT_CHARS:
                            text = text[:MAX_REFERENCE_TEXT_CHARS]
                            job["warnings"].append(
                                f"{item['filename']}: only the first {MAX_REFERENCE_TEXT_CHARS:,} characters were indexed."
                            )
                        material_docs = chunk_text(text, source=item["filename"], method="fixed")
                        for chunk_index, doc in enumerate(material_docs):
                            doc.metadata.update({
                                "source_type": "lecturer_material",
                                "filename": item["filename"],
                                "document_chunk": chunk_index,
                            })
                        docs.extend(material_docs)
                    except Exception as e:
                        warning = f"Could not read {item['filename']}: {friendly_error(e)}"
                        job["warnings"].append(warning)
                        print(f"[job {job_id}] {warning}")

            set_step(job, 3)
            from core.vector_store import (build_vector_store, get_dense_retriever,
                                           new_collection_name, register_lecture)
            from core.bm25_index import build_bm25_retriever
            collection = new_collection_name()
            vector_store = build_vector_store(docs, reset=True, collection_name=collection)
            dense = get_dense_retriever(vector_store, k=10)
            bm25 = build_bm25_retriever(docs, k=10)

            set_step(job, 4)
            import core.reranker  # noqa: F401  (import loads the model)

            set_step(job, 5)
            title = video_title or os.path.splitext(os.path.basename(audio_path))[0]
            lecture_id = register_lecture(title=title, collection_name=collection,
                                          num_chunks=len(docs), transcript_path=transcript_path,
                                          owner_user_id=owner_user_id)
            RUNTIME[lecture_id] = {"docs": docs, "dense": dense, "bm25": bm25}

        job.update(status="done", lecture_id=lecture_id, stage="Workspace ready")
    except Exception as e:
        if lecture_id is None and vector_store is not None:
            try:
                vector_store.delete_collection()
            except Exception:
                pass
        job.update(status="error", error=f"Processing failed during {job.get('stage', 'startup')}. {friendly_error(e)}")
    finally:
        paths_to_remove = {source if not is_url else None, audio_path}
        paths_to_remove.update(item.get("path") for item in reference_files)
        for p in paths_to_remove:
            if p and os.path.exists(p):
                try:
                    os.remove(p)
                except OSError:
                    pass


def new_job() -> str:
    job_id = uuid.uuid4().hex[:12]
    JOBS[job_id] = {"status": "queued", "stage": "Waiting to start", "step": 0,
                    "total": len(STEPS), "lecture_id": None, "error": None, "warnings": []}
    return job_id


def start_job(source: str, is_url: bool, owner_user_id: int, reference_files: list[dict] | None = None):
    job_id = new_job()
    JOBS[job_id]["owner_user_id"] = owner_user_id
    thread = threading.Thread(target=run_job, args=(job_id, source, is_url, owner_user_id, reference_files), daemon=True)
    thread.start()
    return job_id, thread


def job_events(job_id: str, thread: threading.Thread):
    """
    Streams progress as newline-delimited JSON for as long as the job runs.
    The work happens inside this open request on purpose: Cloud Run only gives
    the container CPU while a request is being served.
    """
    last, last_beat = None, time.time()
    while True:
        job = dict(JOBS[job_id])
        alive = thread.is_alive()
        snap = {k: job.get(k) for k in ("status", "stage", "step", "total", "lecture_id", "error", "warnings")}
        if snap != last:
            yield json.dumps(snap) + "\n"
            last, last_beat = snap, time.time()
        elif time.time() - last_beat > 10:
            yield "\n"  # heartbeat so proxies keep the connection open
            last_beat = time.time()
        if job["status"] in ("done", "error"):
            break
        if not alive:
            yield json.dumps({**snap, "status": "error", "error": "Processing stopped unexpectedly."}) + "\n"
            break
        time.sleep(1)


# ── Routes: health + processing ──────────────────────────────────────────────
class RegisterRequest(BaseModel):
    email: str = Field(min_length=5, max_length=254)
    display_name: str = Field(min_length=1, max_length=100)
    password: str = Field(min_length=12, max_length=256)
    role: Literal["student", "farmer"] = "student"


class LoginRequest(BaseModel):
    email: str = Field(min_length=1, max_length=254)
    password: str = Field(min_length=1, max_length=256)


@app.post("/api/auth/register")
def register(req: RegisterRequest, response: Response, request: Request):
    client = request.client.host if request.client else "unknown"
    now = time.time()
    attempts = [t for t in LOGIN_ATTEMPTS.get(client, []) if now - t < LOGIN_RATE_WINDOW_SECONDS]
    LOGIN_ATTEMPTS[client] = attempts
    if len(attempts) >= LOGIN_RATE_LIMIT:
        raise HTTPException(429, "Too many authentication attempts. Try again in a few minutes.")
    attempts.append(now)
    LOGIN_ATTEMPTS[client] = attempts
    try:
        user = auth_store.create_user(req.email, req.display_name, req.password, req.role)
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    token, csrf_token = auth_store.create_session(user["id"])
    set_session_cookie(response, token)
    return {"user": user, "csrf_token": csrf_token}


@app.post("/api/auth/login")
def login(req: LoginRequest, response: Response, request: Request):
    client = request.client.host if request.client else "unknown"
    now = time.time()
    attempts = [t for t in LOGIN_ATTEMPTS.get(client, []) if now - t < LOGIN_RATE_WINDOW_SECONDS]
    if len(attempts) >= LOGIN_RATE_LIMIT:
        LOGIN_ATTEMPTS[client] = attempts
        raise HTTPException(429, "Too many login attempts. Try again in a few minutes.")
    attempts.append(now)
    LOGIN_ATTEMPTS[client] = attempts
    user = auth_store.authenticate_user(req.email, req.password)
    if not user:
        raise HTTPException(401, "Email or password is incorrect, or the account is disabled.")
    LOGIN_ATTEMPTS[client] = []
    token, csrf_token = auth_store.create_session(user["id"])
    set_session_cookie(response, token)
    return {"user": user, "csrf_token": csrf_token}


@app.get("/api/auth/me")
def me(request: Request, user: dict = Depends(current_user)):
    session = auth_store.get_session(request.cookies.get(SESSION_COOKIE))
    return {"user": user, "csrf_token": session["csrf_token"]}


@app.post("/api/auth/logout")
def logout(request: Request, response: Response, user: dict = Depends(current_user)):
    auth_store.revoke_session(request.cookies.get(SESSION_COOKIE))
    clear_session_cookie(response)
    return {"logged_out": True}


@app.get("/api/admin/users")
def admin_users(user: dict = Depends(ADMIN_ONLY)):
    return auth_store.list_users()


class ActiveUserRequest(BaseModel):
    is_active: bool


@app.patch("/api/admin/users/{user_id}/active")
def change_user_active(user_id: int, req: ActiveUserRequest, user: dict = Depends(ADMIN_ONLY)):
    if user_id == user["id"] and not req.is_active:
        raise HTTPException(400, "You cannot disable your own active administrator account.")
    if not auth_store.set_user_active(user_id, req.is_active):
        raise HTTPException(404, "User not found")
    return {"user_id": user_id, "is_active": req.is_active}


@app.get("/api/health")
def health():
    return {
        "status": "ok",
        "max_upload_mb": MAX_UPLOAD_MB,
        "youtube_enabled": YOUTUBE_ENABLED,
        "youtube_captions_enabled": True,
        "max_reference_file_mb": MAX_REFERENCE_FILE_MB,
        "max_reference_files": MAX_REFERENCE_FILES,
        "reference_extensions": sorted(REFERENCE_EXTENSIONS),
    }


def _save_reference_uploads(uploads: list[UploadFile], remaining_bytes: int | None = None) -> list[dict]:
    named_uploads = [item for item in uploads if item and (item.filename or "").strip()]
    if len(named_uploads) > MAX_REFERENCE_FILES:
        raise HTTPException(400, f"Upload at most {MAX_REFERENCE_FILES} lecturer files.")

    paths: list[dict] = []
    max_file_bytes = MAX_REFERENCE_FILE_MB * 1024 * 1024
    budget = MAX_UPLOAD_MB * 1024 * 1024 if remaining_bytes is None else remaining_bytes
    total_bytes = 0
    try:
        for upload_file in named_uploads:
            original_name = os.path.basename(upload_file.filename or "lecturer-material")
            ext = original_name.rsplit(".", 1)[-1].lower() if "." in original_name else ""
            if ext not in REFERENCE_EXTENSIONS:
                accepted = ", ".join(sorted(REFERENCE_EXTENSIONS))
                raise HTTPException(400, f"Unsupported lecturer document '{original_name}'. Use: {accepted}.")

            path = None
            size = 0
            with tempfile.NamedTemporaryFile(delete=False, suffix="." + ext, prefix="lecturer_material_") as tmp:
                path = tmp.name
                while True:
                    chunk = upload_file.file.read(1024 * 1024)
                    if not chunk:
                        break
                    size += len(chunk)
                    total_bytes += len(chunk)
                    if size > max_file_bytes:
                        raise HTTPException(413, f"'{original_name}' is larger than {MAX_REFERENCE_FILE_MB} MB.")
                    if total_bytes > budget:
                        raise HTTPException(413, f"The combined upload exceeds {MAX_UPLOAD_MB} MB. Reduce the video or lecturer documents.")
                    tmp.write(chunk)

            if size == 0:
                os.remove(path)
                raise HTTPException(400, f"'{original_name}' is empty.")
            paths.append({"path": path, "filename": original_name})
        return paths
    except Exception:
        for item in paths:
            try:
                os.remove(item["path"])
            except OSError:
                pass
        if 'path' in locals() and path and os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass
        raise


def extract_reference_text(path: str, filename: str) -> str:
    """Extract text from supported lecturer documents without changing the source file."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        from utils.pdf_reader import read_pdf_text
        text = read_pdf_text(path, with_page_markers=True)
    elif ext in {"txt", "md"}:
        with open(path, "r", encoding="utf-8-sig", errors="replace") as file:
            text = file.read()
    elif ext == "docx":
        from docx import Document as WordDocument
        document = WordDocument(path)
        pieces = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(values):
                    pieces.append(" | ".join(values))
        text = "\n\n".join(pieces)
    elif ext == "pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        slides = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            pieces = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    pieces.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        if any(values):
                            pieces.append(" | ".join(values))
            if pieces:
                slides.append(f"[Slide {slide_number}]\n" + "\n".join(pieces))
        text = "\n\n".join(slides)
    else:
        raise ValueError("Unsupported lecturer document format.")

    text = text.strip()
    if not text:
        raise ValueError("No text could be extracted. Scanned PDFs may need OCR before upload.")
    return text


@app.post("/api/lectures/upload")
def upload(file: UploadFile = File(...), reference_files: list[UploadFile] = File(default=[]), user: dict = Depends(STUDENT_OR_ADMIN)):
    ext = (file.filename or "").rsplit(".", 1)[-1].lower() if "." in (file.filename or "") else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, "Unsupported file type. Use: " + ", ".join(sorted(ALLOWED_EXTENSIONS)))
    size, limit = 0, MAX_UPLOAD_MB * 1024 * 1024
    with tempfile.NamedTemporaryFile(delete=False, suffix="." + ext) as tmp:
        while chunk := file.file.read(1024 * 1024):
            size += len(chunk)
            if size > limit:
                tmp.close()
                os.remove(tmp.name)
                raise HTTPException(413, f"File is larger than {MAX_UPLOAD_MB} MB.")
            tmp.write(chunk)
        path = tmp.name
    if size == 0:
        os.remove(path)
        raise HTTPException(400, "This file is empty.")
    try:
        notes = _save_reference_uploads(reference_files, remaining_bytes=limit - size)
    except Exception:
        try:
            os.remove(path)
        except OSError:
            pass
        raise
    job_id, thread = start_job(path, False, user["id"], notes)
    return StreamingResponse(job_events(job_id, thread), media_type="application/x-ndjson")


class YouTubeRequest(BaseModel):
    url: str


@app.post("/api/lectures/youtube")
async def from_youtube(request: Request, user: dict = Depends(STUDENT_OR_ADMIN)):
    # Keep support for existing JSON clients while accepting multipart form data
    # from the frontend when lecturer documents are attached.
    content_type = request.headers.get("content-type", "").lower()
    if "multipart/form-data" in content_type:
        form = await request.form()
        url = str(form.get("url", "")).strip()
        reference_uploads = form.getlist("reference_files")
    else:
        try:
            payload = await request.json()
        except Exception:
            raise HTTPException(400, "Provide a YouTube URL.")
        url = str(payload.get("url", "")).strip() if isinstance(payload, dict) else ""
        reference_uploads = []

    if not valid_youtube_url(url):
        raise HTTPException(400, "Please provide a link to a single YouTube video.")
    notes = _save_reference_uploads(reference_uploads)
    job_id, thread = start_job(url, True, user["id"], notes)
    return StreamingResponse(job_events(job_id, thread), media_type="application/x-ndjson")


@app.get("/api/jobs/{job_id}")
def job_status(job_id: str, user: dict = Depends(STUDENT_OR_ADMIN)):
    job = JOBS.get(job_id)
    if not job or (user["role"] != "admin" and job.get("owner_user_id") != user["id"]):
        raise HTTPException(404, "Job not found")
    return {k: v for k, v in job.items() if k != "owner_user_id"}


# ── Routes: library (My Lectures) ────────────────────────────────────────────
@app.get("/api/lectures")
def lectures(user: dict = Depends(STUDENT_OR_ADMIN)):
    from core.vector_store import list_lectures
    items = list_lectures()
    if user["role"] != "admin":
        items = [l for l in items if lecture_belongs_to_user(l, user)]
    return [with_title(l) for l in items]


class RenameRequest(BaseModel):
    title: str = Field(min_length=1, max_length=200)


@app.patch("/api/lectures/{lecture_id}")
def rename(lecture_id: str, req: RenameRequest, user: dict = Depends(STUDENT_OR_ADMIN)):
    require_lecture_access(lecture_id, user)
    with TITLE_LOCK:
        titles = load_titles()
        titles[lecture_id] = req.title.strip()
        save_titles(titles)
    return {"id": lecture_id, "title": req.title.strip()}


@app.delete("/api/lectures/{lecture_id}")
def remove(lecture_id: str, user: dict = Depends(STUDENT_OR_ADMIN)):
    from core.vector_store import delete_lecture
    require_lecture_access(lecture_id, user)
    delete_lecture(lecture_id)
    with TITLE_LOCK:
        titles = load_titles()
        if titles.pop(lecture_id, None) is not None:
            save_titles(titles)
    RUNTIME.pop(lecture_id, None)
    SUMMARIES.pop(lecture_id, None)
    DETAILED_NOTES.pop(lecture_id, None)
    for key in [k for k in ENGINES if k[0] == lecture_id]:
        ENGINES.pop(key, None)
    return {"deleted": lecture_id}


# ── Routes: workspace (chat, summary, quiz, evaluation) ──────────────────────
class ChatRequest(BaseModel):
    question: str = Field(min_length=1, max_length=2000)
    session_id: str = Field(default="default", max_length=64)
    response_language: Literal["auto", "English", "Hindi", "Hinglish"] = "auto"
    response_style: Literal["balanced", "simple", "exam", "step_by_step"] = "balanced"
    response_length: Literal["concise", "standard", "detailed"] = "standard"


def source_payload(docs) -> list[dict]:
    return [{"text": d.page_content, "time": d.metadata.get("start_time_str")} for d in docs]


@app.post("/api/lectures/{lecture_id}/chat")
def chat(lecture_id: str, req: ChatRequest, user: dict = Depends(STUDENT_OR_ADMIN)):
    """Streams newline-delimited JSON: one 'sources' line, many 'token' lines, then 'done'."""
    require_lecture_access(lecture_id, user)
    engine = get_engine(lecture_id, req.session_id)
    try:
        stream, sources, save = engine.prepare_stream(
            req.question.strip(),
            preferences={
                "response_language": req.response_language,
                "response_style": req.response_style,
                "response_length": req.response_length,
            },
        )
    except Exception as e:
        raise HTTPException(500, friendly_error(e))

    def events():
        yield json.dumps({"type": "sources", "data": source_payload(sources)}) + "\n"
        parts = []
        try:
            for token in stream:
                parts.append(token)
                yield json.dumps({"type": "token", "data": token}) + "\n"
            save("".join(parts))
            yield json.dumps({"type": "done"}) + "\n"
        except Exception as e:
            yield json.dumps({"type": "error", "data": friendly_error(e)}) + "\n"

    return StreamingResponse(events(), media_type="application/x-ndjson")


@app.post("/api/lectures/{lecture_id}/chat/reset")
def reset_chat(lecture_id: str, session_id: str = "default", user: dict = Depends(STUDENT_OR_ADMIN)):
    require_lecture_access(lecture_id, user)
    ENGINES.pop((lecture_id, session_id), None)
    return {"reset": True}


@app.get("/api/lectures/{lecture_id}/summary")
def summary(lecture_id: str, refresh: bool = False, user: dict = Depends(STUDENT_OR_ADMIN)):
    require_lecture_access(lecture_id, user)
    if refresh or lecture_id not in SUMMARIES:
        from features.summarizer import summarize
        docs = get_runtime(lecture_id)["docs"]
        try:
            SUMMARIES[lecture_id] = summarize(docs)
        except Exception as e:
            raise HTTPException(500, friendly_error(e))
    return {"summary": SUMMARIES[lecture_id]}


@app.get("/api/lectures/{lecture_id}/detailed-notes")
def detailed_notes(lecture_id: str, refresh: bool = False, user: dict = Depends(STUDENT_OR_ADMIN)):
    require_lecture_access(lecture_id, user)
    if refresh or lecture_id not in DETAILED_NOTES:
        from features.detailed_notes import generate_detailed_notes
        docs = get_runtime(lecture_id)["docs"]
        try:
            DETAILED_NOTES[lecture_id] = generate_detailed_notes(docs)
        except Exception as e:
            raise HTTPException(500, friendly_error(e))
    return {"detailed_notes": DETAILED_NOTES[lecture_id]}


class QuizRequest(BaseModel):
    num_questions: int = Field(default=5, ge=1, le=10)


@app.post("/api/lectures/{lecture_id}/quiz")
def quiz(lecture_id: str, req: QuizRequest, user: dict = Depends(STUDENT_OR_ADMIN)):
    require_lecture_access(lecture_id, user)
    from features.quiz_generator import generate_quiz
    docs = get_runtime(lecture_id)["docs"]
    try:
        questions = generate_quiz(docs, num_questions=req.num_questions)
    except Exception as e:
        raise HTTPException(500, friendly_error(e))
    if not questions:
        raise HTTPException(502, "The quiz couldn't be generated this time. Please try again.")
    return {"questions": questions}


class EvalItem(BaseModel):
    question: str
    ground_truth: str


class EvalRequest(BaseModel):
    items: list[EvalItem] = Field(min_length=1, max_length=20)


@app.post("/api/lectures/{lecture_id}/evaluate")
def evaluate(lecture_id: str, req: EvalRequest, user: dict = Depends(ADMIN_ONLY)):
    require_lecture_access(lecture_id, user)
    from features.evaluator import quick_evaluate
    engine = get_engine(lecture_id, "eval-" + uuid.uuid4().hex[:8])  # fresh memory, not the chat's
    items = [i.model_dump() for i in req.items if i.question.strip() and i.ground_truth.strip()]
    if not items:
        raise HTTPException(400, "Each test row needs both a question and an expected answer.")
    try:
        scores = quick_evaluate(engine, items)
    except Exception as e:
        raise HTTPException(500, f"Evaluation failed. {friendly_error(e)}")
    finally:
        for k in [k for k in ENGINES if k[0] == lecture_id and k[1].startswith("eval-")]:
            ENGINES.pop(k, None)
    return {"scores": scores, "n": len(items)}


class FarmCreateRequest(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    location: str = Field(default="", max_length=200)
    crop: str = Field(default="", max_length=120)
    soil_moisture: float | None = Field(default=None, ge=0, le=100)
    temperature_c: float | None = Field(default=None, ge=-50, le=80)
    humidity_pct: float | None = Field(default=None, ge=0, le=100)
    notes: str = Field(default="", max_length=5000)


class FarmUpdateRequest(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=120)
    location: str | None = Field(default=None, max_length=200)
    crop: str | None = Field(default=None, max_length=120)
    soil_moisture: float | None = Field(default=None, ge=0, le=100)
    temperature_c: float | None = Field(default=None, ge=-50, le=80)
    humidity_pct: float | None = Field(default=None, ge=0, le=100)
    notes: str | None = Field(default=None, max_length=5000)


@app.get("/api/farms")
def farms(user: dict = Depends(FARMER_ONLY)):
    return auth_store.list_farms(user["id"])


@app.post("/api/farms")
def create_farm(req: FarmCreateRequest, user: dict = Depends(FARMER_ONLY)):
    data = req.model_dump()
    data["name"] = data["name"].strip()
    if not data["name"]:
        raise HTTPException(400, "Farm name cannot be empty.")
    return auth_store.create_farm(user["id"], data)


@app.get("/api/farms/{farm_id}")
def farm_detail(farm_id: str, user: dict = Depends(FARMER_ONLY)):
    farm = auth_store.get_farm(farm_id)
    if not farm or farm["owner_user_id"] != user["id"]:
        raise HTTPException(404, "Farm not found")
    return farm


@app.patch("/api/farms/{farm_id}")
def edit_farm(farm_id: str, req: FarmUpdateRequest, user: dict = Depends(FARMER_ONLY)):
    current = auth_store.get_farm(farm_id)
    if not current or current["owner_user_id"] != user["id"]:
        raise HTTPException(404, "Farm not found")
    updated = auth_store.update_farm(farm_id, user["id"], req.model_dump(exclude_unset=True))
    if not updated:
        raise HTTPException(404, "Farm not found")
    return updated


@app.delete("/api/farms/{farm_id}")
def remove_farm(farm_id: str, user: dict = Depends(FARMER_ONLY)):
    farm = auth_store.get_farm(farm_id)
    if not farm or farm["owner_user_id"] != user["id"]:
        raise HTTPException(404, "Farm not found")
    if not auth_store.delete_farm(farm_id, user["id"]):
        raise HTTPException(404, "Farm not found")
    return {"deleted": farm_id}


# ── Optional: serve the frontend from the same server (one deployment, no CORS) ─
if os.path.isdir("frontend"):
    app.mount("/", StaticFiles(directory="frontend", html=True), name="frontend")