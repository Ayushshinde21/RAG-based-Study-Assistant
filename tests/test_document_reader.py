import tempfile
import unittest
from pathlib import Path
from unittest import mock

from utils.document_reader import (
    DocumentReadError,
    clean_document_text,
    read_document_text,
)


class TestDocumentTextCleaning(unittest.TestCase):
    def test_normalizes_line_endings_and_removes_null_bytes(self):
        cleaned = clean_document_text("\ufeffTitle\r\nLine 1\x00\r\n\r\n\r\nLine 2  ")
        self.assertEqual(cleaned, "Title\nLine 1\n\nLine 2")

    def test_preserves_markdown_structure_and_source_markers(self):
        source = "# Heading\n\n[Page 4]\n- First point\n- Second point\n\n```python\nx = 1\n```"
        self.assertEqual(clean_document_text(source), source)


class TestDocumentExtraction(unittest.TestCase):
    def test_reads_utf8_markdown_and_preserves_headings(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notes.md"
            path.write_text("\ufeff# RAG\r\n\r\nEmbeddings represent text.\x00", encoding="utf-8")
            text = read_document_text(str(path), "notes.md")

        self.assertEqual(text, "# RAG\n\nEmbeddings represent text.")

    def test_reads_text_file_with_replacement_for_invalid_bytes(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notes.txt"
            path.write_bytes(b"Topic\nValid text\xff")
            text = read_document_text(str(path), "notes.txt")

        self.assertIn("Topic\nValid text", text)
        self.assertIn("\ufffd", text)

    def test_docx_extracts_paragraphs_and_tables(self):
        from docx import Document as WordDocument

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notes.docx"
            document = WordDocument()
            document.add_heading("Retrieval", level=1)
            document.add_paragraph("Embeddings help compare meaning.")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Method"
            table.rows[0].cells[1].text = "Purpose"
            document.save(path)

            text = read_document_text(str(path), "notes.docx")

        self.assertIn("Retrieval", text)
        self.assertIn("Embeddings help compare meaning.", text)
        self.assertIn("Method | Purpose", text)

    def test_pptx_extracts_slide_markers_and_text(self):
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "slides.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            box.text = "Semantic search finds related passages."
            presentation.save(path)

            text = read_document_text(str(path), "slides.pptx")

        self.assertIn("[Slide 1]", text)
        self.assertIn("Semantic search finds related passages.", text)

    def test_pdf_extraction_keeps_page_markers(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notes.pdf"
            path.write_bytes(b"placeholder")
            with mock.patch("utils.pdf_reader.read_pdf_text", return_value="[Page 3]\nPDF text") as reader:
                text = read_document_text(str(path), "notes.pdf")

        reader.assert_called_once_with(str(path), with_page_markers=True)
        self.assertEqual(text, "[Page 3]\nPDF text")

    def test_empty_pdf_extraction_reports_ocr_requirement(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "scan.pdf"
            path.write_bytes(b"placeholder")
            with mock.patch("utils.pdf_reader.read_pdf_text", return_value=""):
                with self.assertRaisesRegex(DocumentReadError, "OCR is required"):
                    read_document_text(str(path), "scan.pdf")

    def test_rejects_unsupported_extension(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "notes.csv"
            path.write_text("a,b", encoding="utf-8")
            with self.assertRaisesRegex(DocumentReadError, "Unsupported document format"):
                read_document_text(str(path), "notes.csv")

    def test_missing_file_raises(self):
        with self.assertRaises(FileNotFoundError):
            read_document_text("missing-document.pdf", "missing-document.pdf")

    def test_empty_text_document_raises(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "empty.txt"
            path.write_text(" \n \n", encoding="utf-8")
            with self.assertRaisesRegex(DocumentReadError, "No readable text"):
                read_document_text(str(path), "empty.txt")


if __name__ == "__main__":
    unittest.main()
